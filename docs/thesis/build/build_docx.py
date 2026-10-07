"""teza_v4.md -> UBT-formatted docx (pandoc + python-docx post-processing)."""
import subprocess, pathlib, copy, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BUILD = pathlib.Path(__file__).parent
SRC_MD = BUILD / "teza_v4.md"
REF = BUILD / "reference.docx"
RAW = BUILD / "raw.docx"
OUT = BUILD / "ANALYTE_Punim_Diplome_Kledion_Berisha.docx"
RES = pathlib.Path("/mnt/user-data/uploads/ANALYTE/docs/thesis")

FONT = "Times New Roman"

def set_font(style, size, bold=None, italic=None, color=None):
    f = style.font
    f.name = FONT
    f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    f.color.rgb = color or RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)

def para_fmt(style, align=None, before=0, after=6, line=1.5, keep_next=None, page_break=None, indent_first=None):
    pf = style.paragraph_format
    if align is not None: pf.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = line
    if keep_next is not None: pf.keep_with_next = keep_next
    if page_break is not None: pf.page_break_before = page_break
    if indent_first is not None: pf.first_line_indent = indent_first

# ---------------- 1. reference docx ----------------
subprocess.run(["pandoc", "-o", str(REF), "--print-default-data-file", "reference.docx"], check=True)
ref = Document(str(REF))
st = ref.styles
class Styles:
    def __init__(self, styles): self.styles = styles
    def __getitem__(self, name):
        for s in self.styles:
            if s.name == name: return s
        raise KeyError(name)
    def add_style(self, *a): return self.styles.add_style(*a)
st = Styles(ref.styles)
def get_or_make(name, base="Normal", kind=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return st[name]
    except KeyError:
        s = st.add_style(name, kind)
        if base: s.base_style = st[base]
        return s

normal = st["Normal"]; set_font(normal, 12); para_fmt(normal, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, 1.5)
for nm in ("Body Text", "First Paragraph", "Compact"):
    s = get_or_make(nm); set_font(s, 12); para_fmt(s, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, 1.5)
h1 = st["Heading 1"]; set_font(h1, 14, True, False); para_fmt(h1, WD_ALIGN_PARAGRAPH.LEFT, 0, 12, 1.5, True, True)
h1.font.all_caps = True
for lvl, size in ((2, 12), (3, 12), (4, 12)):
    h = st[f"Heading {lvl}"]; set_font(h, size, True, lvl == 4); para_fmt(h, WD_ALIGN_PARAGRAPH.LEFT, 12, 6, 1.5, True, False)
for nm in ("Heading 5", "Heading 6", "Title", "Subtitle", "Author", "Date", "Abstract"):
    try:
        s = st[nm]; set_font(s, 12, True); para_fmt(s, WD_ALIGN_PARAGRAPH.LEFT)
    except KeyError:
        pass

fh = get_or_make("FrontHeading"); set_font(fh, 14, True); para_fmt(fh, WD_ALIGN_PARAGRAPH.LEFT, 0, 12, 1.5, True, True); fh.font.all_caps = True
fhn = get_or_make("FrontHeadingNoToc"); set_font(fhn, 14, True); para_fmt(fhn, WD_ALIGN_PARAGRAPH.LEFT, 0, 12, 1.5, True, True); fhn.font.all_caps = True
clg = get_or_make("CoverLogo"); para_fmt(clg, WD_ALIGN_PARAGRAPH.CENTER, 0, 12, 1.0)
cp = get_or_make("CoverProgram"); set_font(cp, 14, True); para_fmt(cp, WD_ALIGN_PARAGRAPH.CENTER, 12, 96, 1.5)
ct = get_or_make("CoverTitle"); set_font(ct, 16, True); para_fmt(ct, WD_ALIGN_PARAGRAPH.CENTER, 24, 72, 1.5)
cl = get_or_make("CoverLine"); set_font(cl, 14, False); para_fmt(cl, WD_ALIGN_PARAGRAPH.CENTER, 0, 36, 1.5)
cf = get_or_make("CoverFoot"); set_font(cf, 12, False, True); para_fmt(cf, WD_ALIGN_PARAGRAPH.CENTER, 72, 0, 1.5)
pb = get_or_make("PageBreak"); set_font(pb, 12); para_fmt(pb, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.0, page_break=True)
tp = get_or_make("TocPlaceholder"); set_font(tp, 12); para_fmt(tp, WD_ALIGN_PARAGRAPH.LEFT, 0, 6, 1.15)
fc = get_or_make("Figure Caption"); set_font(fc, 11, False, True); para_fmt(fc, WD_ALIGN_PARAGRAPH.CENTER, 6, 12, 1.15)
tc = get_or_make("Table Caption"); set_font(tc, 11, True, False); para_fmt(tc, WD_ALIGN_PARAGRAPH.LEFT, 12, 6, 1.15, keep_next=True)
img = get_or_make("Figure"); para_fmt(img, WD_ALIGN_PARAGRAPH.CENTER, 12, 0, 1.0, keep_next=True)
cap = get_or_make("Image Caption"); set_font(cap, 11, False, True); para_fmt(cap, WD_ALIGN_PARAGRAPH.CENTER, 6, 12, 1.15)
for nm in ("Source Code",):
    s = get_or_make(nm); para_fmt(s, WD_ALIGN_PARAGRAPH.LEFT, 0, 6, 1.0)
    s.font.name = "Consolas"; s.font.size = Pt(9)
    rpr = s.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(a), "Consolas")
vc = get_or_make("Verbatim Char", None, WD_STYLE_TYPE.CHARACTER)
vc.font.name = "Consolas"; vc.font.size = Pt(10)
rpr = vc.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
for a in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(a), "Consolas")
bq = get_or_make("Block Text"); set_font(bq, 12, False, True); para_fmt(bq, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, 1.5)
bq.paragraph_format.left_indent = Cm(1)
for nm in ("TOC Heading",):
    try:
        s = st[nm]; set_font(s, 14, True); para_fmt(s, WD_ALIGN_PARAGRAPH.LEFT)
    except KeyError: pass
for i in range(1, 4):
    try:
        s = st[f"TOC {i}"]; set_font(s, 12, i == 1); para_fmt(s, WD_ALIGN_PARAGRAPH.LEFT, 0, 3, 1.15)
        s.paragraph_format.left_indent = Cm(0.6 * (i - 1))
    except KeyError: pass
try:
    s = st["Table of Figures"]; set_font(s, 12); para_fmt(s, WD_ALIGN_PARAGRAPH.LEFT, 0, 3, 1.15)
except KeyError:
    s = st.add_style("Table of Figures", WD_STYLE_TYPE.PARAGRAPH); s.base_style = normal; set_font(s, 12); para_fmt(s, WD_ALIGN_PARAGRAPH.LEFT, 0, 3, 1.15)
# page setup A4, margins 2.5cm
sec = ref.sections[0]
sec.page_width = Cm(21); sec.page_height = Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.5)
def strip_theme_fonts(document):
    root = document.styles.element
    for rf in root.iter(qn("w:rFonts")):
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            if rf.get(qn(a)) is None:
                rf.set(qn(a), FONT)
    for c in root.iter(qn("w:color")):
        c.set(qn("w:val"), "000000")
        for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if c.get(qn(a)) is not None: del c.attrib[qn(a)]
strip_theme_fonts(ref)
ref.save(str(REF))

# ---------------- 2. pandoc ----------------
subprocess.run(["pandoc", str(SRC_MD), "-f", "markdown+pipe_tables+fenced_divs", "-t", "docx",
                "--reference-doc", str(REF), "--resource-path", str(RES) + ":" + str(BUILD), "-o", str(RAW)], check=True)

# ---------------- 3. post-process ----------------
doc = Document(str(RAW))
dst = Styles(doc.styles)
body = doc.element.body

def add_field(paragraph, instr):
    r = paragraph.add_run()
    fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin"); fb.set(qn("w:dirty"), "true")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = " "
    fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end")
    for el in (fb, it, fs, t, fe):
        r._r.append(el)

def para_text(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))

paras = list(doc.paragraphs)
cap_re_fig = re.compile(r"^Figura (\d+|[A-Z]\.\d+)\. ")
cap_re_tab = re.compile(r"^Tabela (\d+|[A-Z]\.\d+)\. ")
for p in paras:
    txt = para_text(p).strip()
    sname = p.style.name
    if txt in ("[[TOC]]", "[[LOF]]", "[[LOT]]"):
        for r in p.runs: r._r.getparent().remove(r._r)
        if txt == "[[TOC]]":
            add_field(p, r' TOC \o "1-3" \h \z \u \t "FrontHeading,1" ')
        elif txt == "[[LOF]]":
            add_field(p, r' TOC \h \z \t "Figure Caption,1" ')
        else:
            add_field(p, r' TOC \h \z \t "Table Caption,1" ')
        continue
    if sname == "PageBreak":
        # ensure it only holds a break
        for r in p.runs: r._r.getparent().remove(r._r)
        continue
    runs = p.runs
    if cap_re_fig.match(txt) and runs and all((r.italic or r.style.name == "Verbatim Char") for r in runs if r.text.strip()):
        p.style = dst["Figure Caption"]
        for r in runs: r.italic = None
        continue
    if cap_re_tab.match(txt) and runs and all((r.italic or r.style.name == "Verbatim Char") for r in runs if r.text.strip()):
        p.style = dst["Table Caption"]
        for r in runs: r.italic = None
        continue
    if sname in ("Figure", "Captioned Figure") or p._p.find(".//" + qn("w:drawing")) is not None:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
    if sname == "Image Caption":
        # pandoc generated caption from alt text ("Figura N") - drop it (we have our own captions)
        p._p.getparent().remove(p._p)

# "LISTA E TABELAVE" must not start a new page
for p in doc.paragraphs:
    if p.style.name == "FrontHeading" and para_text(p).strip() == "LISTA E TABELAVE":
        p.paragraph_format.page_break_before = False

# images: cap width 16 cm, keep aspect
MAXW = Cm(16)
for shp in doc.inline_shapes:
    if shp.width > MAXW:
        ratio = MAXW / shp.width
        shp.height = int(shp.height * ratio); shp.width = MAXW
    if shp.height > Cm(21):
        ratio = Cm(21) / shp.height
        shp.width = int(shp.width * ratio); shp.height = Cm(21)

# tables: borders, font size by width
def set_cell_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}"); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")

for tbl in doc.tables:
    ncols = len(tbl.columns)
    size = 10 if ncols <= 5 else (9 if ncols <= 7 else 8)
    set_cell_borders(tbl)
    tbl.autofit = True
    for ri, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf = p.paragraph_format; pf.space_before = Pt(1); pf.space_after = Pt(1)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for r in p.runs:
                    r.font.size = Pt(size); r.font.name = FONT
                    if ri == 0: r.font.bold = True
    # repeat header row
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    # remove pandoc's "Table" caption-less width 100%
    tblPr = tbl._tbl.tblPr
    w = tblPr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW"); tblPr.append(w)
    w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000")

strip_theme_fonts(doc)

# column widths proportional to content length
from docx.shared import Emu
TEXT_W = Cm(16)
for tbl in doc.tables:
    ncols = len(tbl.columns)
    if ncols < 2: continue
    lens = [1.0] * ncols
    for row in tbl.rows:
        cells = row.cells
        for ci in range(min(ncols, len(cells))):
            L = len(para_text(cells[ci].paragraphs[0]) if cells[ci].paragraphs else "")
            # weight: sqrt-ish to avoid one huge column
            lens[ci] = max(lens[ci], min(L, 80) ** 0.6)
    total = sum(lens)
    widths = [max(Cm(1.9), int(TEXT_W * l / total)) for l in lens]
    scale = TEXT_W / sum(widths)
    widths = [int(w * scale) for w in widths]
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w / 635)))
    for row in tbl.rows:
        for ci, cell in enumerate(row.cells):
            if ci < ncols: cell.width = widths[ci]

# ---------------- sections & page numbers ----------------
def footer_with_page(section, fmt=None, start=None, empty=False):
    section.footer.is_linked_to_previous = False
    ft = section.footer
    for p in list(ft.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p = ft.paragraphs[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if not empty:
        add_field(p, " PAGE ")
        p.runs[0].font.name = FONT; p.runs[0].font.size = Pt(12)
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType"); sectPr.append(pg)
    if fmt: pg.set(qn("w:fmt"), fmt)
    if start is not None: pg.set(qn("w:start"), str(start))

def start_section_before(paragraph):
    """Insert a next-page section break before `paragraph` by ending the previous section
    on a new empty paragraph placed just before it."""
    new_p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    sectPr = copy.deepcopy(doc.sections[-1]._sectPr)
    # strip header/footer refs so sections are independent (we re-add)
    for ref_el in list(sectPr):
        if ref_el.tag in (qn("w:headerReference"), qn("w:footerReference")):
            sectPr.remove(ref_el)
    pPr.append(sectPr); new_p.append(pPr)
    paragraph._p.addprevious(new_p)

abstract_p = next(p for p in doc.paragraphs if p.style.name == "FrontHeadingNoToc" and para_text(p).strip() == "ABSTRAKT")
hyrje_p = next(p for p in doc.paragraphs if p.style.name == "Heading 1" and para_text(p).strip().startswith("1 HYRJE"))
start_section_before(abstract_p)
start_section_before(hyrje_p)
# the heading paragraphs following a section break must not add a page break of their own
abstract_p.paragraph_format.page_break_before = False
hyrje_p.paragraph_format.page_break_before = False

secs = doc.sections
assert len(secs) == 3, len(secs)
for s in secs:
    s.page_width = Cm(21); s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.5)
    s.footer_distance = Cm(1.25)
    s.start_type = WD_SECTION.NEW_PAGE
footer_with_page(secs[0], empty=True)
footer_with_page(secs[1], fmt="upperRoman", start=1)
footer_with_page(secs[2], fmt="decimal", start=1)

# update fields on open
settings = doc.settings.element
uf = settings.find(qn("w:updateFields"))
if uf is None:
    uf = OxmlElement("w:updateFields"); settings.append(uf)
uf.set(qn("w:val"), "true")

# document language: Albanian
for st_ in doc.styles:
    if st_.type == WD_STYLE_TYPE.PARAGRAPH:
        rpr = st_.element.get_or_add_rPr()
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang"); rpr.append(lang)
        lang.set(qn("w:val"), "sq-AL")

doc.core_properties.title = "Bazimi determinist dhe verifikimi i automatizuar i shpjegimeve mjekësore të gjeneruara nga modelet e mëdha gjuhësore"
doc.core_properties.author = "Kledion Berisha"
doc.save(str(OUT))
print("saved", OUT)
